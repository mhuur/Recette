# -*- coding: utf-8 -*-
"""Génère data/saisons/Fruits et légumes de saison.xlsx à partir de data/saisons/saisons.csv.

Valeurs d'une cellule mois dans le CSV :
  2  = pleine saison   → « ● » sur fond foncé
  1  = de saison       → « ● » sur fond clair
  ?  = non vérifié     → « ○ » sur fond jaune pâle
  '' = hors saison / non renseigné → vide

Usage : py tools/saisons_xlsx.py
"""
import csv
import os
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
CSV_PATH = os.path.join(ROOT, "data", "saisons", "saisons.csv")
OUT_PATH = os.path.join(ROOT, "data", "saisons", "Fruits et légumes de saison.xlsx")

MOIS = ["Jan", "Fév", "Mar", "Avr", "Mai", "Juin", "Juil", "Août", "Sep", "Oct", "Nov", "Déc"]

# Couleur par grande famille (clé = début de la catégorie) : (pleine saison, de saison)
PALETTE = {
    "Légume": ("7BB661", "CFE8C0"),
    "Champignon": ("B08968", "E6D5C3"),
    "Herbe": ("4FA3A5", "C9E6E6"),
    "Fruit": ("F2A541", "FBE0B6"),
    "Exotique": ("9E9E9E", "E0E0E0"),
}
FILL_UNK = PatternFill("solid", fgColor="FFF3B0")
FILL_HEAD = PatternFill("solid", fgColor="2F4F3E")
FILL_CAT = PatternFill("solid", fgColor="EEF2EA")
THIN = Side(style="thin", color="D0D0D0")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def palette(cat):
    for key, cols in PALETTE.items():
        if cat.startswith(key):
            return cols
    return ("999999", "DDDDDD")


def main():
    with open(CSV_PATH, encoding="utf-8-sig", newline="") as f:
        rows = list(csv.reader(f, delimiter=";"))
    head, body = rows[0], rows[1:]
    assert head[2:14] == MOIS, head[2:14]

    wb = Workbook()
    ws = wb.active
    ws.title = "Calendrier"

    # En-tête
    for c, name in enumerate(head, start=1):
        cell = ws.cell(row=1, column=c, value=name)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = FILL_HEAD
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER
    ws.row_dimensions[1].height = 30

    r = 2
    last_cat = None
    for row in body:
        cat, name = row[0], row[1]
        months = row[2:14]
        sources = row[14] if len(row) > 14 else ""
        notes = row[15] if len(row) > 15 else ""
        dark, light = palette(cat)

        if cat != last_cat:
            # Ligne de séparation de catégorie
            ws.cell(row=r, column=1, value=cat).font = Font(bold=True)
            for c in range(1, len(head) + 1):
                ws.cell(row=r, column=c).fill = FILL_CAT
                ws.cell(row=r, column=c).border = BORDER
            r += 1
            last_cat = cat

        ws.cell(row=r, column=1, value=cat).font = Font(color="777777", size=9)
        ws.cell(row=r, column=2, value=name)
        for i, v in enumerate(months):
            cell = ws.cell(row=r, column=3 + i)
            cell.alignment = Alignment(horizontal="center", vertical="center")
            v = v.strip()
            if v == "2":
                cell.value = "●"
                cell.fill = PatternFill("solid", fgColor=dark)
                cell.font = Font(color="FFFFFF", bold=True)
            elif v == "1":
                cell.value = "●"
                cell.fill = PatternFill("solid", fgColor=light)
                cell.font = Font(color=dark)
            elif v == "?":
                cell.value = "○"
                cell.fill = FILL_UNK
                cell.font = Font(color="8A6D00")
        ws.cell(row=r, column=15, value=sources).font = Font(size=9, color="555555")
        ws.cell(row=r, column=16, value=notes).font = Font(size=9, color="555555")
        for c in range(1, len(head) + 1):
            ws.cell(row=r, column=c).border = BORDER
        r += 1

    # Totaux par mois
    ws.cell(row=r, column=2, value="Total ● (de saison + pleine)").font = Font(bold=True)
    for i in range(12):
        col = get_column_letter(3 + i)
        cell = ws.cell(row=r, column=3 + i, value=f'=COUNTIF({col}2:{col}{r-1},"●")')
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center")

    # Largeurs, figeage, filtre
    ws.column_dimensions["A"].width = 24
    ws.column_dimensions["B"].width = 52
    for i in range(12):
        ws.column_dimensions[get_column_letter(3 + i)].width = 6
    ws.column_dimensions["O"].width = 60
    ws.column_dimensions["P"].width = 60
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(head))}{r-1}"

    # Feuille Lisez-moi
    lm = wb.create_sheet("Lisez-moi")
    lines = [
        "Fruits et légumes de saison — France métropolitaine",
        "",
        "● fond foncé = pleine saison · ● fond clair = de saison · ○ fond jaune = non vérifié (aucune source consultée) · vide = hors saison ou mois non encore traité.",
        "",
        "État du travail : les 12 mois sont renseignés et sourcés (2026-09-16, tools/saisons_fill.py).",
        "Règle : un mois est « de saison » quand les sources à couverture annuelle qui le listent sont au moins aussi nombreuses que celles qui l'excluent (égalité = oui) ;",
        "« pleine saison » quand, en plus, une source la déclare (pleine saison fruits-legumes.org, cœur de saison Interfel, texte explicite) — cette déclaration pèse double dans le vote.",
        "Les listes de septembre seul (GF, DBB, CV, BM, PG, CDC) ne votent qu'en septembre. « ? » = aucune source pour le produit.",
        "La colonne « Sources (période par source) » donne, pour chaque source, la période qu'elle attribue au produit (pleine saison entre parenthèses) ; la colonne « Notes » signale les réserves.",
        "« = X (assimilé) » : la variété n'est pas distinguée par les sources, elle reprend la valeur du produit X.",
        "",
        "Codes des sources (détail et URL dans data/saisons/SOURCES.md) :",
        "FL  = fruits-legumes.org (calendrier et fiches, pleine saison / de saison)",
        "IF  = Interfel, lesfruitsetlegumesfrais.com (calendrier en ligne et fiches produits : cœur de saison / saison / disponibilité)",
        "GP  = Greenpeace, calendrier « Le guetteur » et guide PDF 2022",
        "BNA = Bio Nouvelle-Aquitaine, calendrier PDF",
        "GF  = Green Fudge, « Les légumes de septembre par famille »",
        "DBB = DocteurBonneBouffe, « Les fruits et légumes de septembre »",
        "CV  = Cerise et Vinaigrette, liste de septembre",
        "BM  = BienManger, calendrier des saisons",
        "PG  = Les Producteurs Gâtinais, calendrier",
        "FM  = La Ferme de Margaux, calendrier de maturité des pommes",
        "ML  = Mangeons local .bzh (courges par variété, aromatiques par mois)",
        "MC  = Mycocarta, saisons des champignons",
        "CDC = Chasseurs de champignons (via résumé de recherche)",
        "SM  = Swissmilk, choux de saison",
        "PDT = lespommesdeterre.com, primeurs et conservation",
        "LN  = LaNutrition.fr, variétés de raisin (via résumé de recherche)",
        "DNC = Drive de nos campagnes, saison de la noix (via résumé de recherche)",
        "VD  = Vedura (via résumé de recherche)",
        "CJ  = Culture Jardin / Aroma-Zone, cébette (via résumé de recherche)",
        "PPF = Pommes et poires de France / ANPP, tableaux variétaux pommes et poires (PDF, périodes de vente)",
        "Sources web des produits sans calendrier (relevé du 2026-09-16, une page lue par produit) : LPF Le Poireau de France · AJ Au Jardin · AGS Agrosemens · POU Pouliquen · JDJ Le Jardin de Jenny ·",
        "JM Jardiner-malin · TV Terre Vivante · IRI Iriso · GB Gerbeaud · BMS Bien manger selon les saisons · RDT La Ratte du Touquet · MDV Mogette de Vendée · ALE Alélor ·",
        "GBQ Graines Bocquet · GSM Graines-semences · ADP Autour du potager · CLC Cultiver-les-champignons · PJ PagesJaunes jardinage/potager · TPM Tisanes et potions des montagnes ·",
        "GRO Pépinières Gromolard · ODC AOP Oignon doux des Cévennes · OTR Office de tourisme de Roscoff · PL pomme-pinklady.com · WK Wikipédia · FDF Fraisiers de France ·",
        "FDS Syndicat AOP Figue de Solliès · PDB Pépinière du Bosc · LAF L'Arbre aux fruits · COC Pépinière Cochet · TLJ Tom le jardinier · MV Monde Végétal · CAN Canaghja ·",
        "AAA Les Agrumes de l'Armagnac · BAK Bakker · PRO Prosem · PNA Produits de Nouvelle-Aquitaine · TD Terre Dauphinoise · IF fiche = fiche produit Interfel (variétés).",
        "",
        "Régénérer le classeur : py tools/saisons_xlsx.py (source : data/saisons/saisons.csv).",
    ]
    for i, line in enumerate(lines, start=1):
        lm.cell(row=i, column=1, value=line)
    lm["A1"].font = Font(bold=True, size=13)
    lm.column_dimensions["A"].width = 140

    wb.save(OUT_PATH)
    print(f"OK {OUT_PATH} — {len(body)} produits")
    return 0


if __name__ == "__main__":
    sys.exit(main())
